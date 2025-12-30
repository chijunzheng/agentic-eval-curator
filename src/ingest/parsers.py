"""Document parsers for various file formats."""

import csv
import json
import re
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

import pdfplumber
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook

from src.models import SourceType


class ParseResult:
    """Result from parsing a document."""

    def __init__(
        self,
        text: str,
        page_texts: list[str] | None = None,
        tables: list[dict[str, Any]] | None = None,
    ):
        """Initialize parse result.

        Args:
            text: Full extracted text.
            page_texts: Optional per-page text list (for PDFs).
            tables: Optional list of extracted tables as dicts.
        """
        self.text = text
        self.page_texts = page_texts
        self.tables = tables or []


class BaseParser(ABC):
    """Abstract base class for document parsers."""

    @abstractmethod
    def parse(self, file_path: Path) -> ParseResult:
        """Parse a document and extract text.

        Args:
            file_path: Path to the document file.

        Returns:
            ParseResult containing extracted text and metadata.
        """
        pass


class TextParser(BaseParser):
    """Parser for plain text and markdown files."""

    def parse(self, file_path: Path) -> ParseResult:
        """Parse a text or markdown file.

        Args:
            file_path: Path to the .txt or .md file.

        Returns:
            ParseResult with file contents.
        """
        text = file_path.read_text(encoding="utf-8")
        return ParseResult(text=text)


class PDFParser(BaseParser):
    """Parser for PDF documents with page reference preservation."""

    def parse(self, file_path: Path) -> ParseResult:
        """Parse a PDF file, extracting text and tables.

        Args:
            file_path: Path to the .pdf file.

        Returns:
            ParseResult with text, per-page text, and tables.
        """
        page_texts: list[str] = []
        tables: list[dict[str, Any]] = []

        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                # Extract text
                page_text = page.extract_text() or ""
                page_texts.append(page_text)

                # Extract tables
                for table_data in page.extract_tables():
                    if table_data:
                        tables.append({
                            "page": page_num,
                            "data": table_data,
                        })

        full_text = "\n\n".join(page_texts)
        return ParseResult(text=full_text, page_texts=page_texts, tables=tables)


class CSVParser(BaseParser):
    """Parser for CSV files with table structure preservation."""

    def parse(self, file_path: Path) -> ParseResult:
        """Parse a CSV file.

        Args:
            file_path: Path to the .csv file.

        Returns:
            ParseResult with text representation and table data.
        """
        rows: list[list[str]] = []

        with open(file_path, newline="", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)

        # Convert to text representation
        text_lines = []
        for row in rows:
            text_lines.append(" | ".join(row))

        # Store as table structure
        tables = [{
            "data": rows,
            "headers": rows[0] if rows else [],
        }]

        return ParseResult(text="\n".join(text_lines), tables=tables)


class JSONParser(BaseParser):
    """Parser for JSON files."""

    def parse(self, file_path: Path) -> ParseResult:
        """Parse a JSON file.

        Args:
            file_path: Path to the .json file.

        Returns:
            ParseResult with flattened text representation.
        """
        with open(file_path, encoding="utf-8") as f:
            data = json.load(f)

        text = self._flatten_json(data)
        return ParseResult(text=text)

    def _flatten_json(self, obj: Any, prefix: str = "") -> str:
        """Recursively flatten JSON to readable text.

        Args:
            obj: JSON object (dict, list, or scalar).
            prefix: Key path prefix for nested values.

        Returns:
            Flattened text representation.
        """
        lines = []

        if isinstance(obj, dict):
            for key, value in obj.items():
                new_prefix = f"{prefix}.{key}" if prefix else key
                lines.append(self._flatten_json(value, new_prefix))
        elif isinstance(obj, list):
            for i, item in enumerate(obj):
                new_prefix = f"{prefix}[{i}]"
                lines.append(self._flatten_json(item, new_prefix))
        else:
            lines.append(f"{prefix}: {obj}")

        return "\n".join(filter(None, lines))


class HTMLParser(BaseParser):
    """Parser for HTML files."""

    def parse(self, file_path: Path) -> ParseResult:
        """Parse an HTML file, stripping tags.

        Args:
            file_path: Path to the .html file.

        Returns:
            ParseResult with extracted text content.
        """
        html_content = file_path.read_text(encoding="utf-8")
        soup = BeautifulSoup(html_content, "html.parser")

        # Remove script and style elements
        for element in soup(["script", "style", "meta", "link"]):
            element.decompose()

        # Get text with whitespace normalization
        text = soup.get_text(separator="\n", strip=True)

        # Extract tables
        tables = []
        for table in soup.find_all("table"):
            rows = []
            for tr in table.find_all("tr"):
                cells = [td.get_text(strip=True) for td in tr.find_all(["td", "th"])]
                if cells:
                    rows.append(cells)
            if rows:
                tables.append({"data": rows})

        return ParseResult(text=text, tables=tables)


class DocxParser(BaseParser):
    """Parser for Microsoft Word (.docx) documents."""

    def parse(self, file_path: Path) -> ParseResult:
        """Parse a DOCX file, extracting text and tables.

        Args:
            file_path: Path to the .docx file.

        Returns:
            ParseResult with extracted text and tables.
        """
        doc = DocxDocument(file_path)

        # Extract text from paragraphs
        paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
        text = "\n\n".join(paragraphs)

        # Extract tables
        tables = []
        for table in doc.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                rows.append(cells)
            if rows:
                tables.append({
                    "data": rows,
                    "headers": rows[0] if rows else [],
                })

        # If there's table data, append text representation
        if tables:
            table_texts = []
            for tbl in tables:
                for row in tbl["data"]:
                    table_texts.append(" | ".join(row))
            if table_texts:
                text = text + "\n\n" + "\n".join(table_texts) if text else "\n".join(table_texts)

        return ParseResult(text=text, tables=tables)


class XlsxParser(BaseParser):
    """Parser for Microsoft Excel (.xlsx) spreadsheets."""

    def parse(self, file_path: Path) -> ParseResult:
        """Parse an XLSX file, extracting data from all sheets.

        Args:
            file_path: Path to the .xlsx file.

        Returns:
            ParseResult with text representation and table data.
        """
        workbook = load_workbook(file_path, data_only=True)

        all_text_lines = []
        tables = []

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]

            # Add sheet name as section header
            all_text_lines.append(f"=== Sheet: {sheet_name} ===")

            rows = []
            for row in sheet.iter_rows(values_only=True):
                # Convert None values to empty strings and stringify all values
                str_row = [str(cell) if cell is not None else "" for cell in row]
                # Skip completely empty rows
                if any(cell.strip() for cell in str_row):
                    rows.append(str_row)
                    all_text_lines.append(" | ".join(str_row))

            if rows:
                tables.append({
                    "sheet": sheet_name,
                    "data": rows,
                    "headers": rows[0] if rows else [],
                })

            all_text_lines.append("")  # Blank line between sheets

        text = "\n".join(all_text_lines).strip()
        return ParseResult(text=text, tables=tables)


class YangParser(BaseParser):
    """Parser for YANG data model files.

    YANG is a data modeling language used for network configuration
    and state data (RFC 6020, RFC 7950). This parser extracts the
    structured content while preserving hierarchy and metadata.
    """

    def parse(self, file_path: Path) -> ParseResult:
        """Parse a YANG file.

        Args:
            file_path: Path to the .yang file.

        Returns:
            ParseResult with YANG model content.
        """
        content = file_path.read_text(encoding="utf-8")

        # Extract metadata from top-level module/submodule
        module_info = self._extract_module_info(content)

        # Build structured text with metadata header
        header_lines = []
        if module_info.get("module"):
            header_lines.append(f"Module: {module_info['module']}")
        if module_info.get("namespace"):
            header_lines.append(f"Namespace: {module_info['namespace']}")
        if module_info.get("prefix"):
            header_lines.append(f"Prefix: {module_info['prefix']}")
        if module_info.get("organization"):
            header_lines.append(f"Organization: {module_info['organization']}")
        if module_info.get("description"):
            header_lines.append(f"Description: {module_info['description']}")

        if header_lines:
            text = "\n".join(header_lines) + "\n\n" + content
        else:
            text = content

        return ParseResult(text=text)

    def _extract_module_info(self, content: str) -> dict[str, str]:
        """Extract key metadata from YANG module.

        Args:
            content: Raw YANG file content.

        Returns:
            Dict with module, namespace, prefix, organization, description.
        """
        info: dict[str, str] = {}

        # Extract module name
        module_match = re.search(r'(?:module|submodule)\s+([^\s{]+)', content)
        if module_match:
            info["module"] = module_match.group(1)

        # Extract namespace
        ns_match = re.search(r'namespace\s+"([^"]+)"', content)
        if ns_match:
            info["namespace"] = ns_match.group(1)

        # Extract prefix
        prefix_match = re.search(r'prefix\s+([^\s;]+)', content)
        if prefix_match:
            info["prefix"] = prefix_match.group(1).strip('"')

        # Extract organization
        org_match = re.search(r'organization\s+"([^"]+)"', content)
        if org_match:
            info["organization"] = org_match.group(1)

        # Extract description (first occurrence, usually module-level)
        desc_match = re.search(r'description\s+"([^"]+)"', content)
        if desc_match:
            info["description"] = desc_match.group(1)

        return info


# Parser registry
_PARSERS: dict[SourceType, type[BaseParser]] = {
    SourceType.TXT: TextParser,
    SourceType.MD: TextParser,
    SourceType.PDF: PDFParser,
    SourceType.CSV: CSVParser,
    SourceType.JSON: JSONParser,
    SourceType.HTML: HTMLParser,
    SourceType.DOCX: DocxParser,
    SourceType.XLSX: XlsxParser,
    SourceType.YANG: YangParser,
}


def get_parser(file_path: Path) -> BaseParser:
    """Get appropriate parser for a file.

    Args:
        file_path: Path to the document.

    Returns:
        Parser instance for the file type.

    Raises:
        ValueError: If file type is not supported.
    """
    extension = file_path.suffix.lower().lstrip(".")
    try:
        source_type = SourceType(extension)
    except ValueError:
        supported = ", ".join(st.value for st in SourceType)
        raise ValueError(
            f"Unsupported file type: {file_path.suffix}. Supported: {supported}"
        )

    parser_class = _PARSERS[source_type]
    return parser_class()
