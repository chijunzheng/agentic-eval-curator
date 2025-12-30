"""Unit tests for document parsers."""

import json
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook

from src.ingest.parsers import (
    CSVParser,
    DocxParser,
    HTMLParser,
    JSONParser,
    TextParser,
    XlsxParser,
    YangParser,
    get_parser,
)


class TestTextParser:
    """Tests for TextParser."""

    def test_parse_txt_file(self, temp_dir: Path):
        """Test parsing a plain text file."""
        file_path = temp_dir / "test.txt"
        content = "Line 1\nLine 2\nLine 3"
        file_path.write_text(content)

        parser = TextParser()
        result = parser.parse(file_path)

        assert result.text == content
        assert result.page_texts is None

    def test_parse_md_file(self, temp_dir: Path):
        """Test parsing a markdown file."""
        file_path = temp_dir / "test.md"
        content = "# Header\n\nParagraph text."
        file_path.write_text(content)

        parser = TextParser()
        result = parser.parse(file_path)

        assert result.text == content

    def test_parse_unicode(self, temp_dir: Path):
        """Test parsing file with unicode characters."""
        file_path = temp_dir / "unicode.txt"
        content = "Hello 世界 🌍"
        file_path.write_text(content, encoding="utf-8")

        parser = TextParser()
        result = parser.parse(file_path)

        assert result.text == content


class TestCSVParser:
    """Tests for CSVParser."""

    def test_parse_csv_simple(self, temp_dir: Path):
        """Test parsing a simple CSV file."""
        file_path = temp_dir / "test.csv"
        content = "name,age,city\nAlice,30,NYC\nBob,25,LA"
        file_path.write_text(content)

        parser = CSVParser()
        result = parser.parse(file_path)

        assert "name | age | city" in result.text
        assert "Alice | 30 | NYC" in result.text
        assert len(result.tables) == 1
        assert result.tables[0]["headers"] == ["name", "age", "city"]

    def test_parse_csv_with_quotes(self, temp_dir: Path):
        """Test parsing CSV with quoted fields."""
        file_path = temp_dir / "quoted.csv"
        content = 'name,desc\n"Smith, John","A description"'
        file_path.write_text(content)

        parser = CSVParser()
        result = parser.parse(file_path)

        assert "Smith, John" in result.text


class TestJSONParser:
    """Tests for JSONParser."""

    def test_parse_json_dict(self, temp_dir: Path):
        """Test parsing a JSON object."""
        file_path = temp_dir / "test.json"
        data = {"name": "Test", "value": 42, "nested": {"key": "val"}}
        file_path.write_text(json.dumps(data))

        parser = JSONParser()
        result = parser.parse(file_path)

        assert "name: Test" in result.text
        assert "value: 42" in result.text
        assert "nested.key: val" in result.text

    def test_parse_json_array(self, temp_dir: Path):
        """Test parsing a JSON array."""
        file_path = temp_dir / "array.json"
        data = [{"id": 1}, {"id": 2}]
        file_path.write_text(json.dumps(data))

        parser = JSONParser()
        result = parser.parse(file_path)

        assert "[0].id: 1" in result.text
        assert "[1].id: 2" in result.text


class TestHTMLParser:
    """Tests for HTMLParser."""

    def test_parse_html_simple(self, temp_dir: Path):
        """Test parsing simple HTML."""
        file_path = temp_dir / "test.html"
        content = """
        <html>
        <head><title>Test</title></head>
        <body>
            <h1>Header</h1>
            <p>Paragraph text.</p>
        </body>
        </html>
        """
        file_path.write_text(content)

        parser = HTMLParser()
        result = parser.parse(file_path)

        assert "Header" in result.text
        assert "Paragraph text." in result.text
        # Script/style should be removed
        assert "<script>" not in result.text

    def test_parse_html_strips_scripts(self, temp_dir: Path):
        """Test that scripts are stripped."""
        file_path = temp_dir / "script.html"
        content = """
        <html>
        <body>
            <p>Text</p>
            <script>alert('evil');</script>
        </body>
        </html>
        """
        file_path.write_text(content)

        parser = HTMLParser()
        result = parser.parse(file_path)

        assert "Text" in result.text
        assert "alert" not in result.text

    def test_parse_html_table(self, temp_dir: Path):
        """Test extracting tables from HTML."""
        file_path = temp_dir / "table.html"
        content = """
        <table>
            <tr><th>A</th><th>B</th></tr>
            <tr><td>1</td><td>2</td></tr>
        </table>
        """
        file_path.write_text(content)

        parser = HTMLParser()
        result = parser.parse(file_path)

        assert len(result.tables) == 1
        assert result.tables[0]["data"][0] == ["A", "B"]
        assert result.tables[0]["data"][1] == ["1", "2"]


class TestDocxParser:
    """Tests for DocxParser."""

    def test_parse_docx_paragraphs(self, temp_dir: Path):
        """Test parsing a DOCX file with paragraphs."""
        file_path = temp_dir / "test.docx"

        # Create a DOCX file
        doc = DocxDocument()
        doc.add_paragraph("First paragraph")
        doc.add_paragraph("Second paragraph")
        doc.save(file_path)

        parser = DocxParser()
        result = parser.parse(file_path)

        assert "First paragraph" in result.text
        assert "Second paragraph" in result.text

    def test_parse_docx_with_table(self, temp_dir: Path):
        """Test parsing a DOCX file with a table."""
        file_path = temp_dir / "table.docx"

        doc = DocxDocument()
        doc.add_paragraph("Intro text")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"
        doc.save(file_path)

        parser = DocxParser()
        result = parser.parse(file_path)

        assert "Intro text" in result.text
        assert len(result.tables) == 1
        assert result.tables[0]["data"][0] == ["A", "B"]
        assert result.tables[0]["data"][1] == ["1", "2"]

    def test_parse_docx_empty(self, temp_dir: Path):
        """Test parsing an empty DOCX file."""
        file_path = temp_dir / "empty.docx"

        doc = DocxDocument()
        doc.save(file_path)

        parser = DocxParser()
        result = parser.parse(file_path)

        assert result.text == ""


class TestXlsxParser:
    """Tests for XlsxParser."""

    def test_parse_xlsx_single_sheet(self, temp_dir: Path):
        """Test parsing an XLSX file with one sheet."""
        file_path = temp_dir / "test.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws["A1"] = "Name"
        ws["B1"] = "Age"
        ws["A2"] = "Alice"
        ws["B2"] = 30
        wb.save(file_path)

        parser = XlsxParser()
        result = parser.parse(file_path)

        assert "=== Sheet: Data ===" in result.text
        assert "Name | Age" in result.text
        assert "Alice | 30" in result.text
        assert len(result.tables) == 1

    def test_parse_xlsx_multiple_sheets(self, temp_dir: Path):
        """Test parsing an XLSX file with multiple sheets."""
        file_path = temp_dir / "multi.xlsx"

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Data1"

        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "Data2"
        wb.save(file_path)

        parser = XlsxParser()
        result = parser.parse(file_path)

        assert "=== Sheet: Sheet1 ===" in result.text
        assert "=== Sheet: Sheet2 ===" in result.text
        assert "Data1" in result.text
        assert "Data2" in result.text
        assert len(result.tables) == 2

    def test_parse_xlsx_skips_empty_rows(self, temp_dir: Path):
        """Test that empty rows are skipped."""
        file_path = temp_dir / "sparse.xlsx"

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Header"
        ws["A3"] = "Value"  # Row 2 is empty
        wb.save(file_path)

        parser = XlsxParser()
        result = parser.parse(file_path)

        assert "Header" in result.text
        assert "Value" in result.text


class TestYangParser:
    """Tests for YangParser."""

    def test_parse_yang_module(self, temp_dir: Path):
        """Test parsing a YANG module file."""
        file_path = temp_dir / "test.yang"
        content = '''module test-module {
    namespace "urn:example:test";
    prefix test;

    organization "Test Org";
    description "A test module";

    container config {
        leaf name {
            type string;
        }
    }
}'''
        file_path.write_text(content)

        parser = YangParser()
        result = parser.parse(file_path)

        assert "Module: test-module" in result.text
        assert "Namespace: urn:example:test" in result.text
        assert "Prefix: test" in result.text
        assert "Organization: Test Org" in result.text
        assert "container config" in result.text

    def test_parse_yang_submodule(self, temp_dir: Path):
        """Test parsing a YANG submodule file."""
        file_path = temp_dir / "sub.yang"
        content = '''submodule test-submodule {
    belongs-to test-module {
        prefix test;
    }

    description "A test submodule";
}'''
        file_path.write_text(content)

        parser = YangParser()
        result = parser.parse(file_path)

        assert "Module: test-submodule" in result.text

    def test_parse_yang_without_metadata(self, temp_dir: Path):
        """Test parsing YANG without standard metadata."""
        file_path = temp_dir / "minimal.yang"
        content = '''// Just some YANG content
container data {
    leaf value {
        type int32;
    }
}'''
        file_path.write_text(content)

        parser = YangParser()
        result = parser.parse(file_path)

        # Should still contain the raw content
        assert "container data" in result.text
        assert "leaf value" in result.text


class TestGetParser:
    """Tests for get_parser factory function."""

    def test_get_parser_txt(self, temp_dir: Path):
        """Test getting parser for .txt file."""
        file_path = temp_dir / "test.txt"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, TextParser)

    def test_get_parser_md(self, temp_dir: Path):
        """Test getting parser for .md file."""
        file_path = temp_dir / "test.md"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, TextParser)

    def test_get_parser_csv(self, temp_dir: Path):
        """Test getting parser for .csv file."""
        file_path = temp_dir / "test.csv"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, CSVParser)

    def test_get_parser_json(self, temp_dir: Path):
        """Test getting parser for .json file."""
        file_path = temp_dir / "test.json"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, JSONParser)

    def test_get_parser_html(self, temp_dir: Path):
        """Test getting parser for .html file."""
        file_path = temp_dir / "test.html"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, HTMLParser)

    def test_get_parser_unsupported(self, temp_dir: Path):
        """Test error for unsupported file type."""
        file_path = temp_dir / "test.xyz"
        file_path.touch()

        with pytest.raises(ValueError, match="Unsupported file type"):
            get_parser(file_path)

    def test_get_parser_docx(self, temp_dir: Path):
        """Test getting parser for .docx file."""
        file_path = temp_dir / "test.docx"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, DocxParser)

    def test_get_parser_xlsx(self, temp_dir: Path):
        """Test getting parser for .xlsx file."""
        file_path = temp_dir / "test.xlsx"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, XlsxParser)

    def test_get_parser_yang(self, temp_dir: Path):
        """Test getting parser for .yang file."""
        file_path = temp_dir / "test.yang"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, YangParser)

    def test_get_parser_case_insensitive(self, temp_dir: Path):
        """Test that file extension matching is case insensitive."""
        file_path = temp_dir / "test.TXT"
        file_path.touch()

        parser = get_parser(file_path)
        assert isinstance(parser, TextParser)
